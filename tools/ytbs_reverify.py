"""Yeni TEİAŞ 3 günlük Excel'i Ağustos tabanı ve mapping ile karşılaştırır.

Kullanım:
  py -3 tools/ytbs_reverify.py --excel "YENİ.xlsx"
"""

from __future__ import annotations

import argparse
import csv
import json
from collections import defaultdict
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import load_workbook

ANLIK_OFFSETS = (6, 12, 18, 24, 30, 36)
SAATLIK_OFFSETS = (7, 13, 19, 25, 31, 37)


def nonempty(value) -> bool:
    if value is None or value == "" or value == "#VALUE!":
        return False
    if isinstance(value, str) and not value.strip():
        return False
    return True


def load_slots(path: Path) -> tuple[list[str], list[dict]]:
    wb = load_workbook(path, data_only=True, read_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    if len(rows) < 3:
        raise SystemExit("Excel beklenen 2 başlık + veri satırına sahip değil.")

    header = rows[0]
    slots = []
    for col in range(4, len(header), 6):
        label = header[col]
        if label:
            slots.append(str(label))

    plants = []
    for row in rows[2:]:
        if row[1] is None:
            continue
        try:
            plant_id = int(row[1])
        except (TypeError, ValueError):
            continue
        anlik = [nonempty(row[c]) if c < len(row) else False for c in ANLIK_OFFSETS]
        saatlik = [nonempty(row[c]) if c < len(row) else False for c in SAATLIK_OFFSETS]
        plants.append(
            {
                "teias_plant_id": plant_id,
                "osb": str(row[0] or "").strip(),
                "plant_name": str(row[2] or "").strip(),
                "any_anlik": any(anlik),
                "any_saatlik": any(saatlik),
                "anlik_slots": anlik,
                "saatlik_slots": saatlik,
            }
        )
    return slots, plants


def pct(part: int, whole: int) -> str:
    if whole == 0:
        return "0%"
    return f"{part / whole:.1%}"


def load_gap_ids(gap_dir: Path, name: str, field: str = "teias_plant_id") -> set[int]:
    path = gap_dir / name
    if not path.exists():
        return set()
    with path.open(encoding="utf-8-sig", newline="") as f:
        return {int(r[field]) for r in csv.DictReader(f) if r.get(field)}


def main() -> None:
    root = Path(__file__).resolve().parents[2]
    parser = argparse.ArgumentParser(description="Yeni TEİAŞ Excel doğrulama")
    parser.add_argument("--excel", required=True)
    parser.add_argument("--out", default=str(root / "ScadaDocs" / "ytbs-gap" / "reverify"))
    parser.add_argument("--gap-dir", default=str(root / "ScadaDocs" / "ytbs-gap"))
    parser.add_argument("--baseline-anlik", type=int, default=90)
    parser.add_argument("--baseline-saatlik", type=int, default=163)
    args = parser.parse_args()

    slots, plants = load_slots(Path(args.excel))
    n = len(plants)
    anlik_n = sum(1 for p in plants if p["any_anlik"])
    saatlik_n = sum(1 for p in plants if p["any_saatlik"])
    both = sum(1 for p in plants if p["any_anlik"] and p["any_saatlik"])
    hourly_only = sum(1 for p in plants if p["any_saatlik"] and not p["any_anlik"])
    neither = sum(1 for p in plants if not p["any_anlik"] and not p["any_saatlik"])

    gap_dir = Path(args.gap_dir)
    unmapped = load_gap_ids(gap_dir, "01_unmapped.csv")
    pasif = load_gap_ids(gap_dir, "02_pasif_only.csv")

    remaining = []
    for p in plants:
        if p["any_anlik"]:
            continue
        reason = []
        if p["teias_plant_id"] in unmapped:
            reason.append("mapping_yok")
        if p["teias_plant_id"] in pasif:
            reason.append("pasif")
        if p["any_saatlik"] and not p["any_anlik"]:
            reason.append("anlik_hala_bos_saatlik_var")
        if not reason:
            reason.append("diger_scada_veya_api")
        remaining.append(
            {
                "teias_plant_id": p["teias_plant_id"],
                "osb": p["osb"],
                "plant_name": p["plant_name"],
                "has_saatlik": p["any_saatlik"],
                "reason": "+".join(reason),
            }
        )

    osb_stats = defaultdict(lambda: {"n": 0, "anlik": 0, "saatlik": 0})
    for p in plants:
        s = osb_stats[p["osb"]]
        s["n"] += 1
        s["anlik"] += int(p["any_anlik"])
        s["saatlik"] += int(p["any_saatlik"])

    slot_rows = []
    for i, label in enumerate(slots):
        a = sum(1 for p in plants if i < len(p["anlik_slots"]) and p["anlik_slots"][i])
        h = sum(1 for p in plants if i < len(p["saatlik_slots"]) and p["saatlik_slots"][i])
        slot_rows.append({"slot": label, "anlik": a, "saatlik": h})

    out = Path(args.out)
    out.mkdir(parents=True, exist_ok=True)
    with (out / "remaining_empty_anlik.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["teias_plant_id", "osb", "plant_name", "has_saatlik", "reason"])
        w.writeheader()
        w.writerows(remaining)

    with (out / "osb_coverage.csv").open("w", encoding="utf-8-sig", newline="") as f:
        w = csv.DictWriter(f, fieldnames=["osb", "plants", "anlik", "saatlik", "anlik_pct", "saatlik_pct"])
        w.writeheader()
        for osb, s in sorted(osb_stats.items(), key=lambda kv: kv[1]["anlik"] / kv[1]["n"]):
            w.writerow(
                {
                    "osb": osb,
                    "plants": s["n"],
                    "anlik": s["anlik"],
                    "saatlik": s["saatlik"],
                    "anlik_pct": pct(s["anlik"], s["n"]),
                    "saatlik_pct": pct(s["saatlik"], s["n"]),
                }
            )

    anlik_ok = anlik_n >= args.baseline_saatlik * 0.9
    gap_closed = hourly_only <= 10
    summary = {
        "generated_at_utc": datetime.now(timezone.utc).isoformat(),
        "excel": str(Path(args.excel).resolve()),
        "plants": n,
        "anlik": anlik_n,
        "saatlik": saatlik_n,
        "both": both,
        "hourly_only": hourly_only,
        "neither": neither,
        "baseline_anlik": args.baseline_anlik,
        "baseline_saatlik": args.baseline_saatlik,
        "anlik_delta_vs_90": anlik_n - args.baseline_anlik,
        "deploy_73_looks_closed": bool(anlik_ok and gap_closed),
        "slots": slot_rows,
        "checks": {
            "anlik_approaches_saatlik": abs(anlik_n - saatlik_n) <= max(8, n * 0.03),
            "anlik_up_from_90": anlik_n > args.baseline_anlik,
            "hourly_only_under_10": hourly_only <= 10,
        },
    }
    (out / "summary.json").write_text(json.dumps(summary, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")

    md = [
        "# YTBS yeniden doğrulama",
        "",
        f"Kaynak: `{Path(args.excel).name}` · {n} santral",
        "",
        f"- Anlık dolu: **{anlik_n}** ({pct(anlik_n, n)}) · Ağustos tabanı {args.baseline_anlik}",
        f"- Saatlik dolu: **{saatlik_n}** ({pct(saatlik_n, n)}) · Ağustos tabanı {args.baseline_saatlik}",
        f"- İkisi de dolu: {both} · yalnızca saatlik: {hourly_only} · ikisi boş: {neither}",
        f"- 73'lük anlık boşluğu kapanmış görünüyor: **{'evet' if summary['deploy_73_looks_closed'] else 'hayır'}**",
        "",
        "## Dilimler",
        "",
        "| Dilim | Anlık | Saatlik |",
        "|---|---:|---:|",
    ]
    for s in slot_rows:
        md.append(f"| {s['slot']} | {s['anlik']} | {s['saatlik']} |")
    md += [
        "",
        "Kalan anlık boşlar: `remaining_empty_anlik.csv` (mapping_yok / pasif / diğer).",
        "OSB kırılımı: `osb_coverage.csv`.",
        "",
    ]
    (out / "YTBS_REVERIFY.md").write_text("\n".join(md), encoding="utf-8")
    print(json.dumps(summary, indent=2, ensure_ascii=False))
    print(f"Rapor: {out / 'YTBS_REVERIFY.md'}")


if __name__ == "__main__":
    main()
